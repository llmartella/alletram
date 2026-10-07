#!/usr/bin/env python3
"""
CALC earnings + near-miss report.

For each request JSON in the input folders:
  1. POST it to /calculations and GET /programs for the same participant/supplier
  2. Join potential_earnings back to every request transaction by transaction_id
  3. Check every transaction against every program it did NOT earn on, and flag:
       - missed incentive date      exact SKU + eligible type, date within N days outside the period
       - different transaction type exact SKU + within dates, but the wrong transaction type
       - product family             first sku_key segment matches + eligible type + within dates
       - eligible, did not earn     exact SKU + type + dates all match, but nothing earned
                                    (usually a qualification, gate or threshold; toggle in config)
     A transaction earns at most once per program, so if it earned on any offer
     in a program, no other offer in that program is checked for it. Otherwise it
     gets at most one line per program on each sheet: missed date shows only the
     closest offer; the other sheets list the program's remaining offers.
     For distributor suppliers (GET /enrollments), transactions whose types need a
     seller (GET /calculations/requirements) must have that distributor as seller;
     a missing or different seller goes straight to No match with the reason.
     Transaction types that never earn (booking, retailer_orders_from_distribution
     by default) are not checked. Gate definitions and zero-rate stacked sku_sets
     never pay, so they never produce flags either.
  4. Write one Excel workbook per retailer (participant_key): an Executive summary,
     a Summary by supplier, program and offer, then one sheet each for Earned,
     Missed date, Product family, Transaction type, Eligible did not earn, and No match

Input layout: one folder per participant_key, one file per program supplier:
  input/{participant_key}/{program_supplier_key}_{time_frame}_<anything>.json
  e.g. input/cpi_hastings_ne_073/basf_2027_invoices_and_purchase_invoices.json
The supplier key is everything before the first _YYYY_, and YYYY is the time_frame.

API calls send program_supplier_key and time_frame as params and the
participant key as the `participant-key` header.

Usage:
  python earnings_report.py                    # call the API for every input file
  python earnings_report.py --reuse-responses  # rebuild the workbook from saved responses, no API calls
  python earnings_report.py --config other.json
"""
import argparse
import json
import re
import os
import sys
from datetime import date, timedelta
from pathlib import Path

import requests

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required for the Excel output. Install it with: pip install openpyxl")

TXN_FIELDS = ["transaction_id", "transaction_type", "sku_key", "uom_key",
              "quantity", "amount", "invoice_date"]
REQUIRED_CONFIG = ("base_url", "input_dir", "output_dir")

FLAG_DATE = "missed incentive date"
FLAG_TYPE = "different transaction type"
FLAG_FAMILY = "product family"
FLAG_ELIGIBLE = "eligible, did not earn"
DEFAULT_NON_EARNING_TYPES = ["booking", "retailer_orders_from_distribution"]


# ---------------------------------------------------------------- config / API

def load_config(path):
    """Read config.json. Lines starting with // are treated as comments."""
    with open(path) as f:
        lines = [ln for ln in f if not ln.lstrip().startswith("//")]
    try:
        cfg = json.loads("".join(lines))
    except json.JSONDecodeError as e:
        sys.exit(f"config.json isn't valid JSON ({e}). Check for a trailing comma "
                 f"before a commented-out line or the closing brace.")
    missing = [k for k in REQUIRED_CONFIG if k not in cfg]
    if missing:
        sys.exit(f"Config is missing: {', '.join(missing)}")
    return cfg


FILENAME_RE = re.compile(r"^(?P<supplier>.+?)_(?P<year>\d{4})(?:_.*)?$")


def parse_filename(path):
    """basf_2027_invoices_and_purchase_invoices.json -> ('basf', '2027')"""
    m = FILENAME_RE.match(path.stem)
    return (m["supplier"], m["year"]) if m else None


def build_headers(cfg, participant):
    """participant key goes in a header; supplier and time_frame go in params."""
    headers = {"Content-Type": "application/json",
               cfg.get("participant_header", "participant-key"): participant}
    token = cfg.get("api_key") or os.environ.get(cfg.get("api_key_env", "CALC_API_KEY"))
    if token:
        headers[cfg.get("auth_header", "Authorization")] = cfg.get("auth_prefix", "Bearer ") + token
    return headers


def call_api(cfg, method, route, participant, supplier, payload=None):
    """supplier=None leaves program_supplier_key out of the params (e.g. /enrollments)."""
    url = cfg["base_url"].rstrip("/") + route
    params = {"time_frame": str(cfg["time_frame"])}
    if supplier:
        params["program_supplier_key"] = supplier
    r = requests.request(method, url, params=params, json=payload,
                         headers=build_headers(cfg, participant),
                         timeout=cfg.get("timeout_seconds", 120))
    if not r.ok:
        raise RuntimeError(f"{method} {route} HTTP {r.status_code}: {r.text[:500]}")
    body = r.json()
    if body.get("errors"):
        print(f"WARN {method} {route}: response errors: {json.dumps(body['errors'])[:500]}")
    return body


_fetched_this_run = set()


def get_or_load(cfg, reuse, cache_path, method, route, participant, supplier, payload=None):
    """Call the API and save the response, or load the saved copy with --reuse-responses.

    A response already fetched earlier in this run (e.g. /enrollments, which is per
    participant) is read from its saved file instead of being fetched again.
    """
    if cache_path in _fetched_this_run:
        return json.loads(cache_path.read_text())
    if reuse:
        if not cache_path.exists():
            raise RuntimeError(f"no saved response at {cache_path}")
        return json.loads(cache_path.read_text())
    body = call_api(cfg, method, route, participant, supplier, payload)
    cache_path.write_text(json.dumps(body, indent=2))
    _fetched_this_run.add(cache_path)
    return body


# ---------------------------------------------------------------- earnings

def flatten_transactions(payload):
    rows = []
    for entity in payload.get("entities", []):
        for t in entity.get("transactions", []):
            row = {"entity_id": entity.get("id", "")}
            row.update({f: t.get(f, "") for f in TXN_FIELDS})
            row["seller_key"] = (t.get("seller") or {}).get("key", "")
            rows.append(row)
    return rows


def earnings_by_transaction(response):
    """{transaction_id: {offer_key: {program_key, offer_key, incentive_amount}}}

    If the same transaction/offer pair appears more than once (e.g. several
    incentive rates), the incentive amounts are summed into one offer entry.
    """
    out = {}
    for block in response.get("data", []):
        for e in block.get("potential_earnings", []):
            tid = str(e.get("transaction_id", ""))
            offer_key = e.get("offer_key", "")
            rec = out.setdefault(tid, {}).setdefault(offer_key, {
                "program_key": e.get("program_key", ""),
                "offer_key": offer_key,
                "incentive_amount": 0.0,
            })
            rec["incentive_amount"] += float(e.get("incentive_amount") or 0)
    return out


def failed_qualifications(calc_response, programs_response):
    """{(program_key, offer_key): [{text, label, actual, target, difference}]} for failed qualifications.

    Waived qualifications count as met. Labels come from the program and offer
    qualifications in GET /programs, falling back to the qualification key.
    """
    labels = {}
    for program in programs_response.get("data", []):
        pk = program.get("key", "")
        for q in program.get("qualifications") or []:
            labels[(pk, q.get("key"))] = q.get("label") or q.get("key")
        for offer in program.get("offers") or []:
            for q in offer.get("qualifications") or []:
                labels[(pk, q.get("key"))] = q.get("label") or q.get("key")

    results = []
    for block in calc_response.get("data", []):
        results.extend(block.get("qualification_results") or [])
    results.extend(calc_response.get("qualification_results") or [])

    notes = {}
    for r in results:
        if r.get("met") or r.get("waived"):
            continue
        pk, ok, qk = r.get("program_key", ""), r.get("offer_key", ""), r.get("qualification_key", "")
        label = labels.get((pk, qk), qk)
        try:
            nums = f" (actual {float(r.get('actual')):,.2f} vs target {float(r.get('target')):,.2f})"
        except (TypeError, ValueError):
            nums = ""
        text = f"Missed: {label}{nums}"
        existing = notes.setdefault((pk, ok), [])
        if all(text != e["text"] for e in existing):
            existing.append({"text": text, "label": label, "actual": r.get("actual"),
                             "target": r.get("target"), "difference": r.get("difference")})
    return notes


# ---------------------------------------------------------------- program rules

def parse_date(value):
    try:
        return date.fromisoformat(str(value)[:10])
    except (TypeError, ValueError):
        return None


def first_segment(sku_key):
    return str(sku_key).split("_", 1)[0]


def offer_rules(offer):
    """Normalize one offer into rules: which SKUs, which transaction type, which period.

    role is "earning" for components that pay (standard/gated incentive rates,
    paying stacked sku_sets, stepped filters), or "gate" / "stack requirement" for
    components that must be present but never pay (gate_definition, zero-rate
    stacked sku_sets). Non-earning rules never produce near-miss flags; a
    transaction that fully matches one is simply expected and left alone.
    """
    strategy = offer.get("strategy_key")
    raw = []  # (period_key, transaction_type, sku_keys, role)

    if strategy in ("standard", "gated"):
        for inc in offer.get("incentives") or []:
            by_period = {}
            for r in inc.get("rates") or []:
                pk = r.get("period_key") or inc.get("period_key")   # gated puts it on the rate
                by_period.setdefault(pk, []).append(r.get("sku_key"))
            for pk, skus in by_period.items():
                raw.append((pk, inc.get("transaction_type"), skus, "earning"))
            gate = inc.get("gate_definition")
            if gate:
                gate_skus = [r.get("sku_key") for r in inc.get("rates") or []]
                raw.append((gate.get("period_key"), gate.get("transaction_type"), gate_skus, "gate"))

    elif strategy == "stacked":
        for s in (offer.get("incentive") or {}).get("sku_sets") or []:
            rates = s.get("rates") or []
            pays = any((r.get("rate") or 0) > 0 for r in rates)
            raw.append((s.get("period_key"), s.get("transaction_type"),
                        [r.get("sku_key") for r in rates],
                        "earning" if pays else "stack requirement"))

    elif strategy == "stepped":
        for f in (offer.get("incentive") or {}).get("filters") or []:
            raw.append((f.get("period_key"), f.get("transaction_type"), f.get("sku_keys") or [], "earning"))

    periods = {p.get("key"): p for p in offer.get("time_periods") or []}
    rules = []
    for pk, ttype, skus, role in raw:
        period = periods.get(pk)
        skus = {s for s in skus if s}
        start = parse_date(period.get("start_date")) if period else None
        end = parse_date(period.get("end_date")) if period else None
        if not (ttype and skus and start and end):
            continue
        rules.append({"period_key": pk, "transaction_type": ttype, "sku_keys": skus,
                      "families": {first_segment(s) for s in skus},
                      "start": start, "end": end, "role": role})
    return rules


def offer_rate_text(offer):
    """Readable earning rate for an offer: "4%", "$25.00/unit", "0.5%-2% (stepped)"."""
    strategy = offer.get("strategy_key")
    rates = []
    if strategy in ("standard", "gated"):
        for inc in offer.get("incentives") or []:
            rates += inc.get("rates") or []
    elif strategy == "stacked":
        for sset in (offer.get("incentive") or {}).get("sku_sets") or []:
            rates += [r for r in sset.get("rates") or [] if (r.get("rate") or 0) > 0]
    elif strategy == "stepped":
        rates += (offer.get("incentive") or {}).get("rates") or []

    def fmt(rate, measure):
        if measure == "percent":
            return f"{rate:g}%"
        if measure == "dollars":
            return f"${rate:,.2f}/unit"
        return f"{rate:g} {measure or ''}".strip()

    by_measure = {}
    for r in rates:
        if r.get("rate") is not None:
            by_measure.setdefault(r.get("measure"), set()).add(float(r["rate"]))
    parts = []
    for measure, vals in by_measure.items():
        lo, hi = min(vals), max(vals)
        parts.append(fmt(lo, measure) if lo == hi else f"{fmt(lo, measure)}-{fmt(hi, measure)}")
    text = ", ".join(parts)
    if text and strategy == "stepped":
        text += " (stepped)"
    elif text and strategy == "gated":
        text += " (up to ordered qty)"
    return text


def load_offers(programs_response):
    """[{program_key, offer_key, strategy, qualifier, rules}], plus warnings for unreadable offers."""
    offers, warnings = [], []
    for program in programs_response.get("data", []):
        for offer in program.get("offers") or []:
            rules = offer_rules(offer)
            if not rules:
                warnings.append(f"{offer.get('key')} (strategy '{offer.get('strategy_key')}'): "
                                f"couldn't read SKUs/types/dates, skipped for near-miss checks")
                continue
            offers.append({
                "program_key": program.get("key", ""),
                "program_label": program.get("program_name_label") or program.get("key", ""),
                "offer_key": offer.get("key", ""),
                "strategy": offer.get("strategy_key", ""),
                "rate_text": offer_rate_text(offer),
                "qualifier": offer.get("qualifier") or program.get("qualifier") or "",
                "rules": rules,
            })
    return offers, warnings


# ---------------------------------------------------------------- near-miss checks

def check_rule(txn, rule, window_days):
    """Return None, {"flag": "match"}, or a review dict with the flag and its details."""
    d = parse_date(txn["invoice_date"])
    sku, ttype = txn["sku_key"], txn["transaction_type"]
    exact = sku in rule["sku_keys"]
    family = not exact and first_segment(sku) in rule["families"]
    type_ok = ttype == rule["transaction_type"]
    start, end = rule["start"], rule["end"]
    in_dates = d is not None and start <= d <= end
    base = {"period_key": rule["period_key"], "period_start": start, "period_end": end}

    if exact and type_ok and in_dates:
        return {"flag": "match"}
    if exact and type_ok and d is not None:
        window = timedelta(days=window_days)
        if start - window <= d < start:
            return {**base, "flag": FLAG_DATE, "days_off": (start - d).days, "direction": "before start"}
        if end < d <= end + window:
            return {**base, "flag": FLAG_DATE, "days_off": (d - end).days, "direction": "after end"}
        return None
    if exact and not type_ok and in_dates:
        return {**base, "flag": FLAG_TYPE, "required_type": rule["transaction_type"]}
    if family and type_ok and in_dates:
        similar = [s for s in rule["sku_keys"] if first_segment(s) == first_segment(sku)]
        # closest = longest shared start of the key, then alphabetical
        similar.sort(key=lambda s: (-len(os.path.commonprefix([s, sku])), s))
        return {**base, "flag": FLAG_FAMILY, "closest_sku": similar[0], "other_skus": similar[1:]}
    return None


def review_transaction(txn, earned, offers, cfg):
    """All review dicts for one transaction, one per offer + flag type.

    A transaction earns at most once per program, so once it has earned on any
    offer in a program, the other offers in that program are not checked. The
    same goes for a program where it fully matches an offer (right SKU, type and
    dates) but didn't earn, usually a missed qualification: it belongs to that
    offer, so near misses on the program's other offers would be misleading.
    """
    earned_programs = {o["program_key"] for o in earned}
    # Bookings and orders never earn on their own (they only satisfy a gate or a
    # zero-rate stack requirement), so near-miss flags don't apply to them.
    non_earning = set(cfg.get("non_earning_transaction_types", DEFAULT_NON_EARNING_TYPES))
    if txn["transaction_type"] in non_earning:
        return []
    window = cfg.get("near_miss_days", 30)
    flag_eligible = cfg.get("flag_eligible_not_earned", True)

    # Pass 1: check every offer, and note programs where an earning rule fully matched.
    checked, matched_programs = [], set()
    for offer in offers:
        if offer["program_key"] in earned_programs:
            continue
        results = [check_rule(txn, rule, window) for rule in offer["rules"]]
        matched = [rule for rule, res in zip(offer["rules"], results) if res and res["flag"] == "match"]
        if any(r["role"] == "earning" for r in matched):
            matched_programs.add(offer["program_key"])
        checked.append((offer, results, matched))

    # Pass 2: build review lines.
    reviews = []
    for offer, results, matched in checked:
        ident = {"program_key": offer["program_key"], "offer_key": offer["offer_key"],
                 "program_label": offer.get("program_label") or offer["program_key"]}
        if matched:
            # Fully matches this offer: expected for gate/stack pieces; otherwise worth a look.
            earning = [r for r in matched if r["role"] == "earning"]
            if flag_eligible and earning:
                reviews.append({**ident, "flag": FLAG_ELIGIBLE, "qualifier": offer["qualifier"],
                                "rate_text": offer.get("rate_text", ""),
                                "period_key": earning[0]["period_key"],
                                "period_start": earning[0]["start"], "period_end": earning[0]["end"]})
            continue
        if offer["program_key"] in matched_programs:
            continue   # belongs to another offer in this program; not a near miss here
        # Only earning components produce near-miss flags. Gate definitions and
        # zero-rate stack requirements never pay, so they only suppress flags above.
        seen = set()
        for rule, res in zip(offer["rules"], results):
            if rule["role"] != "earning":
                continue
            if res and res["flag"] not in seen:
                seen.add(res["flag"])
                reviews.append({**ident, **res})
    return collapse_per_program(reviews)


def collapse_per_program(reviews):
    """At most one line per program and flag type.

    Missed date keeps only the offer the transaction came closest to (staggered
    offers: Jan-Feb, Mar-Apr, May-Jun). Other flags keep the first offer by key and
    list the program's remaining offers with the same miss in other_offers.
    """
    grouped = {}
    for rv in reviews:
        grouped.setdefault((rv["program_key"], rv["flag"]), []).append(rv)
    out = []
    for (_, flag), items in grouped.items():
        if flag == FLAG_DATE:
            best = min(items, key=lambda rv: (rv["days_off"], rv["offer_key"]))
            out.append({**best, "other_offers": []})
        else:
            items.sort(key=lambda rv: rv["offer_key"])
            out.append({**items[0], "other_offers": [rv["offer_key"] for rv in items[1:]]})
    return out


# ---------------------------------------------------------------- seller rule

def supplier_type(enrollments, supplier):
    """program_supplier_type for this supplier from GET /enrollments, or None if not listed."""
    for e in (enrollments or {}).get("data", []):
        if e.get("program_supplier_key") == supplier:
            return e.get("program_supplier_type")
    return None


def seller_transaction_types(requirements):
    """Transaction types that require seller.key, from GET /calculations/requirements."""
    types = set()
    for block in (requirements or {}).get("data", []):
        for f in ((block.get("transaction_requirements") or {}).get("fields")) or []:
            if f.get("field") == "seller":
                types.update(f.get("transaction_types") or [])
    return types


def seller_problem(txn, supplier, seller_types):
    """For distributor programs: the seller must be the distributor itself."""
    if txn["transaction_type"] not in seller_types:
        return None
    seller = txn.get("seller_key") or ""
    if not seller:
        return ("Missing seller", f"{txn['transaction_type']} needs seller {supplier}")
    if seller != supplier:
        return ("Different seller", f"seller is {seller}; {supplier} programs need seller {supplier}")
    return None


# ---------------------------------------------------------------- no-match reasons

def no_match_reason(txn, offers, cfg):
    """Why a transaction didn't earn and wasn't a near miss: (reason, detail)."""
    ttype, sku = txn["transaction_type"], txn["sku_key"]
    if ttype in set(cfg.get("non_earning_transaction_types", DEFAULT_NON_EARNING_TYPES)):
        return ("Transaction type never earns", f"{ttype} only counts toward gates or stacks")
    rules = [r for o in offers for r in o["rules"] if r["role"] == "earning"]
    if not any(sku in r["sku_keys"] for r in rules):
        if any(first_segment(sku) in r["families"] for r in rules):
            return ("Product family only, type or date also off",
                    f"'{first_segment(sku)}' products are in a program, but not this SKU")
        return ("SKU not in any program", "")
    if any(check_rule(txn, r, 0) == {"flag": "match"} for r in rules):
        return ("Matched an offer but did not earn", "eligible flag is turned off in config")
    if any(sku in r["sku_keys"] and ttype == r["transaction_type"] for r in rules):
        days = cfg.get("near_miss_days", 30)
        return (f"Date more than {days} days outside incentive period", "")
    return ("Transaction type and date both off", "")


# ---------------------------------------------------------------- records

def build_records(txns, earnings, offers, cfg, supplier, seller_types, qual_notes=None):
    """[(transaction, earned offers, review dicts)] for every request transaction.

    Transactions with no earnings and no reviews get txn["no_match_reason"] and
    txn["no_match_detail"] for the No match sheet.
    """
    records = []
    for t in txns:
        earned = list(earnings.get(str(t["transaction_id"]), {}).values())
        problem = None if earned else seller_problem(t, supplier, seller_types)
        reviews = [] if problem else review_transaction(t, earned, offers, cfg)
        for rv in reviews:
            fails = (qual_notes or {}).get((rv["program_key"], rv["offer_key"]), [])
            rv["qualification_failures"] = fails
            rv["qualification_note"] = "; ".join(f["text"] for f in fails)
        if not earned and not reviews:
            t["no_match_reason"], t["no_match_detail"] = problem or no_match_reason(t, offers, cfg)
        records.append((t, earned, reviews))
    return records


# ---------------------------------------------------------------- Excel workbook

FONT_NAME = "Arial"
F_BASE = Font(name=FONT_NAME, size=10)
F_BOLD = Font(name=FONT_NAME, size=10, bold=True)
F_SHEET_TITLE = Font(name=FONT_NAME, size=14, bold=True)
F_SECTION = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
F_NOTE = Font(name=FONT_NAME, size=10, italic=True, color="595959")
F_GROUP = Font(name=FONT_NAME, size=10, bold=True, color="305496")
FILL_SECTION = PatternFill("solid", fgColor="305496")
FILL_HEADER = PatternFill("solid", fgColor="D9E1F2")
FILL_TOTAL = PatternFill("solid", fgColor="F2F2F2")
FILL_PROGRAM = PatternFill("solid", fgColor="E2EFDA")
FMT_MONEY = '$#,##0.00;($#,##0.00);"-"'
FMT_COUNT = '#,##0;(#,##0);"-"'
FMT_QTY = '#,##0.##'
FMT_DATE = "yyyy-mm-dd"

# (header, transaction field, number format)
TXN_COLS = [("Transaction ID", "transaction_id", None),
            ("Transaction type", "transaction_type", None),
            ("SKU", "sku_key", None),
            ("UOM", "uom_key", None),
            ("Quantity", "quantity", FMT_QTY),
            ("Amount", "amount", FMT_MONEY),
            ("Invoice date", "invoice_date", FMT_DATE),
            ("Seller", "seller_key", None)]


def _num(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return value


def _txn_values(txn):
    vals = []
    for _, field, fmt in TXN_COLS:
        v = txn.get(field, "")
        if fmt == FMT_DATE:
            v = parse_date(v) or v
        elif fmt in (FMT_QTY, FMT_MONEY):
            v = _num(v)
        vals.append(v)
    return vals


def _put(ws, row, col, value, font=F_BASE, fmt=None, fill=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = font
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    return cell


def _section_title(ws, row, text, width):
    for c in range(1, width + 1):
        _put(ws, row, c, None, fill=FILL_SECTION)
    _put(ws, row, 1, text, font=F_SECTION, fill=FILL_SECTION)
    return row + 1


def _header(ws, row, headers):
    for c, h in enumerate(headers, start=1):
        _put(ws, row, c, h, font=F_BOLD, fill=FILL_HEADER)
    return row + 1


def _rows(ws, row, values_list, formats):
    for values in values_list:
        for c, (v, fmt) in enumerate(zip(values, formats), start=1):
            _put(ws, row, c, v, fmt=fmt)
        row += 1
    return row


def _autosize(ws, max_width=60):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            # skip titles, section bars (larger fonts) and italic notes; they can overflow
            if cell.value is None or cell.font.i or (cell.font.sz or 10) > 10:
                continue
            v = cell.value
            n = 10 if hasattr(v, "year") else len(str(v))
            widths[cell.column] = max(widths.get(cell.column, 0), n)
    for col, n in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(n + 2, 10), max_width)


def _sheet_name(name, used):
    clean = "".join(ch for ch in name if ch not in '[]:*?/\\')[:31] or "Sheet"
    base, i = clean, 2
    while clean.lower() in used:
        suffix = f" ({i})"
        clean = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(clean.lower())
    return clean


def _finish_table(ws, header_row, last_row, ncols):
    """Freeze below the header and switch on filters for the whole table."""
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ncols)}{max(last_row, header_row)}"
    _autosize(ws)


def _write_earned_sheet(ws, participant, suppliers):
    """One continuous, filterable table: one line per earned transaction + offer."""
    txn_headers = [h for h, _, _ in TXN_COLS]
    txn_formats = [f for _, _, f in TXN_COLS]
    headers = ["Program supplier", "Program", "Offer", *txn_headers, "Earned amount"]
    formats = [None, None, None, *txn_formats, FMT_MONEY]
    earn_col = len(headers)
    L = get_column_letter(earn_col)

    lines = []
    for supplier in sorted(suppliers):
        for txn, earned, _ in suppliers[supplier]:
            for o in earned:
                lines.append([supplier, o["program_key"], o["offer_key"], *_txn_values(txn),
                              round(o["incentive_amount"], 2)])
    lines.sort(key=lambda r: (r[0], r[1], r[2], str(r[3])))

    _put(ws, 1, 1, f"Earned | {participant}", font=F_SHEET_TITLE)
    header_row = 4
    first, last = header_row + 1, header_row + max(len(lines), 1)
    _put(ws, 2, earn_col - 1, "Total earned (filtered rows)", font=F_BOLD)
    _put(ws, 2, earn_col, f"=SUBTOTAL(9,{L}{first}:{L}{last})", font=F_BOLD, fmt=FMT_MONEY)
    _header(ws, header_row, headers)
    end = _rows(ws, first, lines, formats) - 1
    _finish_table(ws, header_row, end, len(headers))


def _write_review_sheet(ws, title, participant, suppliers, flag, extra_headers, extra_formats, extra_values):
    """One near-miss reason as a continuous, filterable table."""
    txn_headers = [h for h, _, _ in TXN_COLS]
    txn_formats = [f for _, _, f in TXN_COLS]
    headers = ["Program supplier", "Program", "Offer", *txn_headers, *extra_headers]
    formats = [None, None, None, *txn_formats, *extra_formats]

    lines = []
    for supplier in sorted(suppliers):
        for txn, _, reviews in suppliers[supplier]:
            for rv in reviews:
                if rv["flag"] == flag:
                    lines.append([supplier, rv["program_key"], rv["offer_key"], *_txn_values(txn),
                                  *extra_values(rv)])
    lines.sort(key=lambda r: (r[0], r[1], r[2], str(r[3])))

    _put(ws, 1, 1, f"{title} | {participant} ({len(lines)})", font=F_SHEET_TITLE)
    header_row = 3
    _header(ws, header_row, headers)
    end = _rows(ws, header_row + 1, lines, formats) - 1
    _finish_table(ws, header_row, end, len(headers))


def _write_no_match_sheet(ws, participant, suppliers):
    """Transactions that didn't earn and weren't a near miss, with the reason."""
    txn_headers = [h for h, _, _ in TXN_COLS]
    txn_formats = [f for _, _, f in TXN_COLS]
    headers = ["Program supplier", *txn_headers, "Reason", "Detail"]
    formats = [None, *txn_formats, None, None]

    lines = []
    for supplier in sorted(suppliers):
        for t, earned, reviews in suppliers[supplier]:
            if not earned and not reviews:
                lines.append([supplier, *_txn_values(t), t.get("no_match_reason", ""),
                              t.get("no_match_detail", "")])
    reason_col = len(headers) - 2
    lines.sort(key=lambda r: (r[0], r[reason_col], str(r[1])))

    _put(ws, 1, 1, f"No match | {participant} ({len(lines)})", font=F_SHEET_TITLE)
    header_row = 3
    _header(ws, header_row, headers)
    end = _rows(ws, header_row + 1, lines, formats) - 1
    _finish_table(ws, header_row, end, len(headers))


REVIEW_SHEETS = [
    # (sheet name, flag, extra headers, extra formats, extra values)
    ("Missed date", FLAG_DATE,
     ["Incentive start", "Incentive end", "Missed by (days)", "Before / after", "Qualification"],
     [FMT_DATE, FMT_DATE, FMT_COUNT, None, None],
     lambda rv: [rv["period_start"], rv["period_end"], rv["days_off"], rv["direction"],
                 rv.get("qualification_note", "")]),
    ("Product family", FLAG_FAMILY,
     ["Closest offer SKU", "Other matching offer SKUs", "Qualification", "Other offers"],
     [None, None, None, None],
     lambda rv: [rv["closest_sku"], ", ".join(rv["other_skus"]), rv.get("qualification_note", ""),
                 ", ".join(rv.get("other_offers", []))]),
    ("Transaction type", FLAG_TYPE,
     ["Earning transaction type", "Incentive start", "Incentive end", "Qualification", "Other offers"],
     [None, FMT_DATE, FMT_DATE, None, None],
     lambda rv: [rv["required_type"], rv["period_start"], rv["period_end"],
                 rv.get("qualification_note", ""), ", ".join(rv.get("other_offers", []))]),
    ("Eligible, did not earn", FLAG_ELIGIBLE,
     ["Qualification", "Incentive start", "Incentive end"],
     [None, FMT_DATE, FMT_DATE],
     lambda rv: [rv.get("qualification_note") or "No failed qualification found; check gates/thresholds",
                 rv["period_start"], rv["period_end"]]),
]


def _new_bucket():
    return {"earning": set(), "earnings": 0.0, "date": set(), "family": set(), "type": set(), "eligible": set()}


def _summary_rows(records):
    """{program_key: {"all": bucket, "offers": {offer_key: bucket}}}, counting distinct transactions."""
    progs = {}
    flag_field = {FLAG_DATE: "date", FLAG_FAMILY: "family", FLAG_TYPE: "type", FLAG_ELIGIBLE: "eligible"}

    def buckets(program, offer):
        p = progs.setdefault(program, {"all": _new_bucket(), "offers": {}})
        return p["all"], p["offers"].setdefault(offer, _new_bucket())

    for txn, earned, reviews in records:
        tid = str(txn["transaction_id"])
        for o in earned:
            for b in buckets(o["program_key"], o["offer_key"]):
                b["earning"].add(tid)
                b["earnings"] += o["incentive_amount"]
        for rv in reviews:
            for b in buckets(rv["program_key"], rv["offer_key"]):
                b[flag_field[rv["flag"]]].add(tid)
    return progs


def _bucket_values(b):
    return [len(b["earning"]), round(b["earnings"], 2), len(b["date"]),
            len(b["family"]), len(b["type"]), len(b["eligible"])]


def _amount(txn):
    try:
        return float(txn.get("amount") or 0)
    except (TypeError, ValueError):
        return 0.0


def _exec_section(ws, row, title, note, headers, formats, lines):
    width = len(headers)
    row = _section_title(ws, row, title, width)
    _put(ws, row, 1, note, font=F_NOTE)
    row = _header(ws, row + 1, headers)
    if not lines:
        _put(ws, row, 1, "None")
        return row + 2
    return _rows(ws, row, lines, formats) + 1


SIDE_GROUP = Side(style="thin", color="8EA9DB")
FILL_PAIR = PatternFill("solid", fgColor="F3F6FB")
CENTER = Alignment(horizontal="center", vertical="center")


def _date_section(ws, row, buckets, lines):
    """Missed-date table with a three-level header: Early/Late, bucket, then Txns / $.

    Each bucket's two columns share a merged header, a left border, and
    alternating light shading so the count and dollars read as one group.
    """
    n_pairs = 2 * len(buckets) + 1           # early buckets, late buckets, total
    width = 2 + 2 * n_pairs
    row = _section_title(ws, row, "Missed the incentive dates", width)
    _put(ws, row, 1, "Right product and transaction type, but invoiced just before the program started "
                     "(early) or just after it ended (late).", font=F_NOTE)
    top, mid, sub = row + 1, row + 2, row + 3

    def head(r, c, value, span=1):
        for cc in range(c, c + span):
            cell = _put(ws, r, cc, None, font=F_BOLD, fill=FILL_HEADER)
            cell.alignment = CENTER
        _put(ws, r, c, value, font=F_BOLD, fill=FILL_HEADER).alignment = CENTER
        if span > 1:
            ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + span - 1)

    for c, name in ((1, "Program supplier"), (2, "Program")):
        head(top, c, name)
        ws.merge_cells(start_row=top, start_column=c, end_row=sub, end_column=c)
        ws.cell(row=top, column=c).alignment = Alignment(vertical="bottom")

    pair_cols = []                            # first column of each pair
    col = 3
    for title in ("Early (before start)", "Late (after end)"):
        head(top, col, title, span=2 * len(buckets))
        for _, _, label in buckets:
            head(mid, col, label, span=2)
            head(sub, col, "Txns")
            head(sub, col + 1, "$")
            pair_cols.append(col)
            col += 2
    # "Total" spans two rows and two columns: style every cell, then merge once
    # (merging the same cells twice makes Excel report the file as damaged).
    for r in (top, mid):
        for cc in (col, col + 1):
            _put(ws, r, cc, None, font=F_BOLD, fill=FILL_HEADER).alignment = CENTER
    _put(ws, top, col, "Total", font=F_BOLD, fill=FILL_HEADER).alignment = CENTER
    ws.merge_cells(start_row=top, start_column=col, end_row=mid, end_column=col + 1)
    head(sub, col, "Txns")
    head(sub, col + 1, "$")
    pair_cols.append(col)

    row = sub + 1
    if not lines:
        _put(ws, row, 1, "None")
        row += 1
    else:
        formats = [None, None] + [FMT_COUNT, FMT_MONEY] * n_pairs
        row = _rows(ws, row, lines, formats)

    for i, c in enumerate(pair_cols):
        shade = FILL_PAIR if i % 2 == 0 else None
        for r in range(mid, row):
            for cc in (c, c + 1):
                cell = ws.cell(row=r, column=cc)
                if r > sub and shade:
                    cell.fill = shade
            ws.cell(row=r, column=c).border = Border(left=SIDE_GROUP)
    for r in range(top, row):
        ws.cell(row=r, column=pair_cols[0]).border = Border(left=Side(style="medium", color="305496"))
        ws.cell(row=r, column=pair_cols[len(buckets)]).border = Border(left=Side(style="medium", color="305496"))
        ws.cell(row=r, column=pair_cols[-1]).border = Border(left=Side(style="medium", color="305496"))
    return row + 1


def _write_executive_sheet(ws, participant, suppliers, cfg):
    window = cfg.get("near_miss_days", 30)
    # Upper edge of each bucket; the last bucket runs to the near-miss window.
    edges = [e for e in cfg.get("exec_date_bucket_edges", [7, 14, 24]) if e < window] + [window]
    buckets, low = [], 1
    for e in edges:
        buckets.append((low, e, f"{low}-{e} days"))
        low = e + 1

    def bucket_of(days):
        return next(label for lo, hi, label in buckets if lo <= days <= hi)
    _put(ws, 1, 1, f"Executive summary | {participant}", font=F_SHEET_TITLE)
    _put(ws, 2, 1, f"time_frame {cfg['time_frame']}. Amounts are net of returns. Missed date, product family "
                   f"and transaction type leave out anything that also missed a qualification.", font=F_NOTE)
    row = 4

    def reviews(flag, qualified):
        """(supplier, review, txn) for a flag; qualified=True drops failed qualifications, False keeps only them."""
        for supplier in sorted(suppliers):
            for txn, _, rvs in suppliers[supplier]:
                for rv in rvs:
                    if rv["flag"] == flag and bool(rv.get("qualification_failures")) != qualified:
                        yield supplier, rv, txn

    # ---- 1. Missed date, split early / late and by distance
    cols = [(d, label) for d in ("before start", "after end") for _, _, label in buckets]
    labels = {}  # program_key -> program name label (rows are keyed by key, shown by label)
    agg = {}
    for supplier, rv, txn in reviews(FLAG_DATE, qualified=True):
        key = (supplier, rv["program_key"])
        labels[rv["program_key"]] = rv["program_label"]
        a = agg.setdefault(key, {c: [0, 0.0] for c in cols})
        c = (rv["direction"], bucket_of(rv["days_off"]))
        a[c][0] += 1
        a[c][1] += _amount(txn)
    lines = []
    for (supplier, program), a in sorted(agg.items(), key=lambda kv: (kv[0][0], labels[kv[0][1]], kv[0][1])):
        vals = []
        for c in cols:
            vals += [a[c][0], round(a[c][1], 2)]
        lines.append([supplier, labels[program], *vals,
                      sum(a[c][0] for c in cols), round(sum(a[c][1] for c in cols), 2)])
    row = _date_section(ws, row, buckets, lines)
    date_end = row

    # ---- 2. Product family
    agg = {}
    for supplier, rv, txn in reviews(FLAG_FAMILY, qualified=True):
        labels[rv["program_key"]] = rv["program_label"]
        key = (supplier, rv["program_key"], txn["sku_key"], rv["closest_sku"])
        a = agg.setdefault(key, [0, 0.0])
        a[0] += 1
        a[1] += _amount(txn)
    lines = sorted([[k[0], labels[k[1]], *k[2:], n, round(amt, 2)] for k, (n, amt) in agg.items()],
                   key=lambda r: tuple(str(x) for x in r))
    row = _exec_section(ws, row, "Similar products that aren't in the program",
                        "You bought a product from the same family as one the program pays on, but this exact "
                        "SKU isn't configured to earn. Worth confirming if you expected it to.",
                        ["Program supplier", "Program", "SKU purchased", "Closest SKU the program pays on",
                         "Transactions", "Amount ($)"],
                        [None, None, None, None, FMT_COUNT, FMT_MONEY], lines)

    # ---- 3. Transaction type
    agg = {}
    for supplier, rv, txn in reviews(FLAG_TYPE, qualified=True):
        labels[rv["program_key"]] = rv["program_label"]
        key = (supplier, rv["program_key"], txn["transaction_type"], rv["required_type"])
        a = agg.setdefault(key, [0, 0.0])
        a[0] += 1
        a[1] += _amount(txn)
    lines = sorted([[k[0], labels[k[1]], *k[2:], n, round(amt, 2)] for k, (n, amt) in agg.items()],
                   key=lambda r: tuple(str(x) for x in r))
    row = _exec_section(ws, row, "Right product and dates, different transaction type",
                        "These transactions were for program products within the program dates, but were "
                        "reported as a transaction type the program doesn't pay on.",
                        ["Program supplier", "Program", "Transaction type reported", "Program pays on",
                         "Transactions", "Amount ($)"],
                        [None, None, None, None, FMT_COUNT, FMT_MONEY], lines)

    # ---- 4. Eligible, waiting on a qualification
    agg = {}
    for supplier, rv, txn in reviews(FLAG_ELIGIBLE, qualified=False):
        for f in rv["qualification_failures"]:
            key = (supplier, rv["program_key"], rv["offer_key"], f["label"])
            labels[rv["program_key"]] = rv["program_label"]
            a = agg.setdefault(key, {"n": 0, "amt": 0.0, "rate": rv.get("rate_text", ""), "f": f})
            a["n"] += 1
            a["amt"] += _amount(txn)
    lines = []
    for (supplier, program, offer, label), a in sorted(agg.items(), key=lambda kv: (kv[0][0], labels[kv[0][1]], kv[0][2], kv[0][3])):
        f = a["f"]
        try:
            needed = abs(float(f["difference"])) if f.get("difference") is not None \
                else float(f["target"]) - float(f["actual"])
        except (TypeError, ValueError):
            needed = None
        lines.append([supplier, labels[program], offer, label, a["n"], round(a["amt"], 2), a["rate"],
                      f.get("actual"), f.get("target"), needed])
    _exec_section(ws, row, "Ready to earn once a qualification is met",
                  "Right product, dates and transaction type. These should start earning once the "
                  "qualification is met. Actual, target and still needed are in the qualification's own "
                  "units (dollars, units, percent or index). An offer missing two qualifications shows on two rows.",
                  ["Program supplier", "Program", "Offer", "Qualification missed", "Transactions",
                   "Amount ($)", "Program rate", "Actual", "Target", "Still needed"],
                  [None, None, None, None, FMT_COUNT, FMT_MONEY, None, "#,##0.00", "#,##0.00", "#,##0.00"],
                  lines)
    # Columns A-B fit their content; C onward share one width so the Txns / $ pairs
    # line up, and long text in the lower sections wraps instead of widening them.
    _autosize(ws)
    for c in range(3, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(c)].width = 20
    wrap = Alignment(wrap_text=True, vertical="top")
    for r in range(date_end, ws.max_row + 1):
        for c in range(3, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            if isinstance(cell.value, str) and cell.font.sz in (None, 10) and not cell.font.i:
                cell.alignment = wrap
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def write_workbook(path, participant, suppliers, cfg):
    wb = Workbook()
    _write_executive_sheet(wb.active, participant, suppliers, cfg)
    wb.active.title = "Executive summary"
    ws = wb.create_sheet("Summary")

    ALL = "All offers"
    headers = ["Program supplier", "Program", "Offer", "Earning transactions", "Total earnings",
               "Missed date only", "Product family only", "Transaction type only",
               "Eligible, did not earn"]
    formats = [None, None, None, FMT_COUNT, FMT_MONEY, FMT_COUNT, FMT_COUNT, FMT_COUNT, FMT_COUNT]
    first_num = 4  # first numeric column
    _put(ws, 1, 1, f"Summary | {participant}", font=F_SHEET_TITLE)
    _put(ws, 2, 1, f"time_frame {cfg['time_frame']}. Counts are distinct transactions. "
                   f"\"{ALL}\" rows count each transaction once per program, so they can be lower "
                   f"than the sum of the offer rows below them. Totals add up the \"{ALL}\" rows only.")
    row = _header(ws, 4, headers)
    table_first = row
    offer_col = get_column_letter(3)

    def total_row(r, label, first, last, font=F_BOLD, fill=None):
        for c in range(1, len(headers) + 1):
            _put(ws, r, c, None, fill=fill)
        _put(ws, r, 2, label, font=font, fill=fill)
        for c in range(first_num, len(headers) + 1):
            col = get_column_letter(c)
            _put(ws, r, c, f'=SUMIFS({col}{first}:{col}{last},{offer_col}{first}:{offer_col}{last},"{ALL}")',
                 font=font, fmt=formats[c - 1], fill=fill)

    for supplier in sorted(suppliers):
        progs = _summary_rows(suppliers[supplier])
        first = row
        _put(ws, row, 1, supplier, font=F_GROUP)
        row += 1
        if not progs:
            _put(ws, row, 2, "No earnings or near misses")
            row += 1
        for program in sorted(progs):
            p = progs[program]
            # program row (bold), then its offers indented beneath, collapsible
            for c, (v, fmt) in enumerate(zip([None, program, ALL, *_bucket_values(p["all"])], formats), start=1):
                _put(ws, row, c, v, font=F_BOLD, fmt=fmt, fill=FILL_PROGRAM if c > 1 else None)
            row += 1
            offers_first = row
            for offer in sorted(p["offers"]):
                row = _rows(ws, row, [[None, None, offer, *_bucket_values(p["offers"][offer])]], formats)
            if row > offers_first:
                ws.row_dimensions.group(offers_first, row - 1, outline_level=1, hidden=False)
        total_row(row, f"{supplier} total", first, row - 1, fill=FILL_TOTAL)
        row += 2

    total_row(row, "Grand total", table_first, row - 1)
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.freeze_panes = ws.cell(row=table_first, column=1)
    _autosize(ws)

    _write_earned_sheet(wb.create_sheet("Earned"), participant, suppliers)
    for name, flag, headers_x, formats_x, values_x in REVIEW_SHEETS:
        _write_review_sheet(wb.create_sheet(name), name, participant, suppliers,
                            flag, headers_x, formats_x, values_x)
    _write_no_match_sheet(wb.create_sheet("No match"), participant, suppliers)

    wb.save(path)


# ---------------------------------------------------------------- main

def process_file(cfg, path, participant, out_dir, reuse):
    """Call the APIs for one input file and return (supplier, time_frame, records)."""
    keys = parse_filename(path)
    if not keys:
        print(f"SKIP {participant}/{path.name}: expected {{program_supplier_key}}_{{YYYY}}_....json")
        return None
    supplier, time_frame = keys
    cfg = {**cfg, "time_frame": time_frame}     # the file's year drives every API call
    label = f"{participant}/{path.name}"

    payload = json.loads(path.read_text())
    if payload.get("participant_key") and payload["participant_key"] != participant:
        print(f"WARN {label}: participant_key in file is '{payload['participant_key']}'")

    try:
        calc = get_or_load(cfg, reuse, out_dir / "responses" / participant / path.name,
                           "POST", "/calculations", participant, supplier, payload)
        programs = get_or_load(cfg, reuse, out_dir / "programs" / participant / path.name,
                               "GET", "/programs", participant, supplier)
    except Exception as e:
        print(f"FAIL {label}: {e}")
        return None

    # Seller rule: only for distributor suppliers, only on the transaction types
    # /calculations/requirements says need a seller. Failures here don't stop the run.
    seller_types = set()
    try:
        enrollments = get_or_load(cfg, reuse, out_dir / "enrollments" / f"{participant}_{time_frame}.json",
                                  "GET", "/enrollments", participant, None)
        stype = supplier_type(enrollments, supplier)
        if stype is None:
            print(f"WARN {label}: '{supplier}' isn't in /enrollments for {participant}; seller check skipped")
        elif stype == "distributor":
            requirements = get_or_load(cfg, reuse, out_dir / "requirements" / participant / path.name,
                                       "GET", "/calculations/requirements", participant, supplier)
            seller_types = seller_transaction_types(requirements)
    except Exception as e:
        print(f"WARN {label}: seller check skipped ({e})")

    txns = flatten_transactions(payload)
    earnings = earnings_by_transaction(calc)
    offers, warnings = load_offers(programs)
    for w in warnings:
        print(f"WARN {label}: {w}")

    orphans = set(earnings) - {str(t["transaction_id"]) for t in txns}
    if orphans:
        print(f"WARN {label}: {len(orphans)} earning transaction_id(s) not in request, "
              f"e.g. {sorted(orphans)[:5]}")

    records = build_records(txns, earnings, offers, cfg, supplier, seller_types,
                            failed_qualifications(calc, programs))
    earned = sum(1 for _, e, _ in records if e)
    flagged = sum(1 for _, _, r in records if r)
    print(f"OK   {label}: {len(records)} transactions, {earned} earned, {flagged} flagged for review, "
          f"{len(offers)} offers checked")
    return supplier, time_frame, records


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--config", default="config.json")
    ap.add_argument("--reuse-responses", action="store_true",
                    help="rebuild the workbook from saved API responses instead of calling the API")
    args = ap.parse_args()

    # Resolve paths relative to the script (for config.json) and to the config
    # file (for input/output), so it works no matter which folder you run it from.
    config_path = Path(args.config)
    if not config_path.is_absolute():
        config_path = Path(__file__).resolve().parent / config_path
    cfg = load_config(config_path)
    base = config_path.parent
    in_dir = base / cfg["input_dir"]
    out_dir = base / cfg["output_dir"]
    folders = sorted(d for d in in_dir.iterdir() if d.is_dir()) if in_dir.is_dir() else []
    if not folders:
        sys.exit(f"No participant folders in {in_dir}. Expected input/{{participant_key}}/{{supplier}}_{{YYYY}}_....json")

    for participant_dir in folders:
        participant = participant_dir.name
        files = sorted(participant_dir.glob("*.json"))
        if not files:
            print(f"SKIP {participant}: no .json files")
            continue
        for sub in ("responses", "programs", "requirements"):
            (out_dir / sub / participant).mkdir(parents=True, exist_ok=True)
        (out_dir / "enrollments").mkdir(parents=True, exist_ok=True)

        # One file per supplier per folder; if there are more, skip that supplier
        # rather than guess which file is the right one.
        by_supplier = {}
        for path in files:
            keys = parse_filename(path)
            by_supplier.setdefault(keys[0] if keys else None, []).append(path)
        dupes = {k: v for k, v in by_supplier.items() if k and len(v) > 1}
        for supplier, paths in dupes.items():
            print(f"SKIP {participant}: {len(paths)} files for '{supplier}' "
                  f"({', '.join(p.name for p in paths)}); keep one and run again")

        suppliers, years = {}, set()
        for path in files:
            keys = parse_filename(path)
            if keys and keys[0] in dupes:
                continue
            result = process_file(cfg, path, participant, out_dir, args.reuse_responses)
            if result:
                supplier, time_frame, records = result
                suppliers[supplier] = records
                years.add(time_frame)
        if not suppliers:
            continue
        if len(years) > 1:
            print(f"WARN {participant}: files cover more than one year ({', '.join(sorted(years))})")
        wb_path = out_dir / f"{participant}.xlsx"
        write_workbook(wb_path, participant, suppliers, {**cfg, "time_frame": ", ".join(sorted(years))})
        print(f"XLSX {participant}: {len(suppliers)} supplier(s), 8 sheets -> {wb_path}")


if __name__ == "__main__":
    main()